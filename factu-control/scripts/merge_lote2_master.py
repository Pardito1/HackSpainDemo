"""Copia el maestro y le añade los proveedores y pedidos del lote 2.

El Excel original no se toca: se escribe un libro nuevo. El upsert va por clave
(`ID` en Proveedores, `Pedido` en Pedidos_2026) y dice en voz alta cuántas filas
añadió y cuántas actualizó, porque `version_fuentes` depende del resultado.
"""

import argparse
import csv
import hashlib
import json
import shutil
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from factu.columns import normalized  # noqa: E402


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    return [{normalized(k): v for k, v in row.items()} for row in rows]


def number(value):
    """Un importe del CSV entra como número; lo demás se queda como texto."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def upsert(sheet, rows, key, numeric):
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    columns = [normalized(name) for name in header]
    if normalized(key) not in columns:
        raise ValueError(f"{sheet.title}: falta la columna clave {key}")
    key_index = columns.index(normalized(key))
    existing = {
        str(row[key_index].value).strip(): row
        for row in sheet.iter_rows(min_row=2)
        if row[key_index].value is not None
    }
    added, updated = [], []
    for row in rows:
        identifier = str(row.get(normalized(key), "")).strip()
        if not identifier:
            continue
        values = [
            number(row[c]) if c in numeric else row.get(c)
            for c in columns
        ]
        if identifier in existing:
            for cell, value in zip(existing[identifier], values):
                if value is not None:
                    cell.value = value
            updated.append(identifier)
        else:
            sheet.append(values)
            added.append(identifier)
    return added, updated


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True)
    parser.add_argument("--proveedores", required=True)
    parser.add_argument("--pedidos", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    source, output = Path(args.excel), Path(args.output)
    if source.resolve() == output.resolve():
        parser.error("El maestro original no se modifica: elige otro --output")
    shutil.copyfile(source, output)

    book = openpyxl.load_workbook(output)
    report = {}
    for sheet, path, key, numeric in (
        ("Proveedores", args.proveedores, "ID", set()),
        ("Pedidos_2026", args.pedidos, "Pedido", {"importetotal"}),
    ):
        added, updated = upsert(book[sheet], read_csv(path), key, numeric)
        report[sheet] = {
            "añadidas": len(added),
            "actualizadas": len(updated),
            "filas": book[sheet].max_row - 1,
            "nuevas": added,
        }
    book.save(output)
    book.close()

    report["output"] = str(output)
    report["sha256"] = hashlib.sha256(output.read_bytes()).hexdigest()
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
