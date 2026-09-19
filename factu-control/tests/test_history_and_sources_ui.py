import copy
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from factu.historical import history_dependency
from factu.master import read_master
from factu.policy import evaluate
from factu.source_view import impact_counts
from factu.service import Service
from factu.web import create_app


def add_archive(path, rows):
    book = openpyxl.load_workbook(path)
    sheet = book.create_sheet("Pedidos_2025_OLD")
    sheet.append(["Pedido", "Importe"])
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()


def processed(bundle):
    service, batch, _, _ = bundle
    service.process(batch, ocr=False)
    return service, batch


def test_archive_parses_records_not_notes_or_formulas(bundle):
    _, _, _, path = bundle
    add_archive(path, [("PO-2025-0812", 4100.5), ("PO-2025-0977", 233.1),
                       ("(archivo parcial, resto en backup_marzo??)", None),
                       ("=HYPERLINK(\"https://invalid\")", 100), ("PO-2025-0100", "=1/0")])
    original = path.read_bytes()
    archive = read_master(path)["historical"]
    assert len(archive["orders"]) == 3 and len(archive["notes"]) == 2
    assert archive["orders"]["PO-2025-0812"][0]["total"] == "4100.5"
    assert archive["orders"]["PO-2025-0812"][0]["source"]["cells"] == {"order": "A2", "total": "B2"}
    assert archive["orders"]["PO-2025-0100"][0]["total"] is None
    assert len(archive["warnings"]) == 1 and archive["coverage"] == "partial"
    assert path.read_bytes() == original


def test_history_matching_requires_order_not_amount(facts):
    extraction, master, erp, policy = facts
    master["historical"] = {"available": True, "orders": {"PO-2025-0001": [{"id":"PO-2025-0001", "total":"121.00", "source":{"sheet":"Pedidos_2025_OLD", "row":2}}]}}
    assert evaluate(*facts, "2026-09-19")["result"] == "PAGAR"  # Same suffix/amount is not a duplicate.
    master["historical"]["orders"]["PO-2026-0001"] = master["historical"]["orders"].pop("PO-2025-0001")
    decision = evaluate(*facts, "2026-09-19")
    assert decision["result"] == "ESCALAR"
    rule = next(r for r in decision["rules"] if r["id"] == "historical_order")
    assert rule["state"] == "FAIL" and rule["evidence"]["rows"][0]["source"]["row"] == 2
    # Independent confirmed payment can still justify not paying; the old sheet cannot.
    erp["rows"][0]["estado"] = "PAGADA"
    assert evaluate(*facts, "2026-09-19")["result"] == "NO_PAGAR"


def test_legacy_snapshot_cannot_claim_history_checked(facts):
    facts[1].pop("historical")
    facts[1]["sheets"].append("Pedidos_2025_OLD")
    decision = evaluate(*facts, "2026-09-19")
    assert decision["result"] == "ESCALAR"
    assert next(r for r in decision["rules"] if r["id"] == "historical_order")["state"] == "UNKNOWN"


def test_history_change_is_selective_audited_and_does_not_repeat_ocr(bundle):
    service, batch = processed(bundle)
    doc = service.store.one("SELECT * FROM documents")
    old = service.store.source(service.batch(batch)["master_id"])
    master = copy.deepcopy(old)
    master["historical"] = {"available": True, "coverage": "partial", "orders": {}}
    plan = service.preview_source(batch, "master", service.store.put_source("master", master), "Ana", "Archivo histórico incorporado")
    service.commit_source(plan["id"])
    before = service.document(doc["id"])["latest_decision"]
    master["historical"]["orders"]["PO-2025-0999"] = [{"id":"PO-2025-0999", "total":"121.00", "source":{"sheet":"Pedidos_2025_OLD","row":2}}]
    unrelated = service.preview_source(batch, "master", service.store.put_source("master", master), "", "")
    assert len(unrelated["reused"]) == 1 and not unrelated["affected"]
    service.commit_source(unrelated["id"])
    assert service.document(doc["id"])["latest_decision"] == before
    master["historical"]["orders"]["PO-2026-0001"] = [{"id":"PO-2026-0001", "total":"121.00", "source":{"sheet":"Pedidos_2025_OLD","row":3}}]
    related = service.preview_source(batch, "master", service.store.put_source("master", master), "", "")
    assert related["affected"][0]["after"] == "ESCALAR"
    assert service.export_rows(batch)[0]["result"] == "PAGAR"
    service.commit_source(related["id"])
    assert service.export_rows(batch)[0]["result"] == "ESCALAR"
    assert service.document(doc["id"])["extraction"] == doc["extraction"]
    assert service.store.verify_audit()["valid"]


def test_source_preview_uses_indexed_duplicate_check(bundle, monkeypatch):
    service, batch = processed(bundle)
    doc = service.store.one("SELECT * FROM documents")
    expected = service.duplicate_context(doc, service.effective(doc))
    original = service.duplicate_context
    calls = []
    def indexed(document, extraction, candidates=None):
        assert candidates is not None
        calls.append(document["id"])
        result = original(document, extraction, candidates=candidates)
        assert result == expected
        return result
    monkeypatch.setattr(service, "duplicate_context", indexed)
    master = service.store.source(service.batch(batch)["master_id"])
    master["historical"] = {"available":True, "orders": {"PO-2026-0001":[{"id":"PO-2026-0001"}]}}
    plan = service.preview_source(batch, "master", service.store.put_source("master", master), "", "")
    assert calls == [doc["id"]] and plan["affected"][0]["after"] == "ESCALAR"


def test_historical_possible_duplicate_needs_explicit_human_evidence(bundle):
    service, batch = processed(bundle)
    doc = service.store.one("SELECT * FROM documents")
    master = service.store.source(service.batch(batch)["master_id"])
    master["historical"] = {"available": True, "orders": {"PO-2026-0001": [{"id":"PO-2026-0001","total":"121.00"}]}}
    plan = service.preview_source(batch, "master", service.store.put_source("master", master), "Ana", "Nuevo antecedente")
    service.commit_source(plan["id"])
    args = (doc["id"], "PAGAR", "Alberto", "Contabilidad confirma que son obligaciones diferentes.", "Confirmación de contabilidad expediente 34")
    with TestClient(create_app(service.store.root)) as client:
        html = client.get(f"/documents/{doc['id']}").text
        assert 'name="acknowledged" value="historical_order"' in html
        assert "El pedido aparece en el archivo de años anteriores" in html
    with pytest.raises(ValueError, match="expresamente"):
        service.human_preview(*args, [])
    preview = service.human_preview(*args, ["historical_order"])
    service.human_commit(*args, ["historical_order"], 0, preview["preview_token"])
    assert service.detail(doc["id"])["decision"]["engine_result"] == "ESCALAR"
    assert service.export_rows(batch)[0]["result"] == "PAGAR"


def test_impact_counts_are_disjoint():
    plan = {"affected": [{"before":a,"after":b} for a,b in [("PAGAR","PAGAR"),("ESCALAR","ESCALAR"),("PAGAR","ESCALAR"),("PAGAR","NO_PAGAR")]], "reused":[{}]*3}
    assert impact_counts(plan) == {"affected":4,"unchanged":2,"to_review":1,"other_changes":1,"reused":3}


def test_first_erp_query_is_not_presented_as_update_preview(bundle, tmp_path):
    _, _, folder, workbook = bundle
    service = Service(tmp_path / "initial-query")
    batch = service.ingest(folder, workbook, "Primer lote", "2026-09-19")["batch_id"]
    with TestClient(create_app(service.store.root)) as client:
        html = client.get(f"/sources?batch={batch}").text
        assert "Consultar el ERP por primera vez" in html
        assert 'data-action="sync"' in html
        assert 'class="erp-change-form"' not in html
        assert 'class="source-form"' not in html


def test_source_ui_optional_notes_no_fake_identity_and_preview_before_apply(bundle):
    service, batch = processed(bundle)
    _, _, _, workbook = bundle
    add_archive(workbook, [("PO-2026-0001",121)])
    app = create_app(service.store.root)
    with TestClient(app) as client:
        headers = {"X-CSRF-Token": app.state.service.store.csrf_token}
        response = client.post(f"/api/batches/{batch}/changes/upload", data={"kind":"master"}, files={"file":("updated.xlsx",workbook.read_bytes())}, headers=headers)
        assert response.status_code == 200, response.text
        plan = response.json()
        assert plan["actor"] == "Sesión local (sin identificar)"
        assert service.export_rows(batch)[0]["result"] == "PAGAR"
        html = client.get(f"/sources?batch={batch}&change={plan['id']}").text
        for text in ["Datos y actualizaciones", "Ver facturas afectadas", "Datos contables (ERP)", "Revisar cambios", "Aplicar actualización", "Pasan a revisión", "Historial de actualizaciones"]:
            assert text in html
        for forbidden in ["SHA-256", "/api/sources/", "Makefile", "Quién revisa el cambio", "Qué ha cambiado", "Demo sintética"]:
            assert forbidden not in html
        assert 'name="reason" required' not in html
        assert client.post(f"/api/changes/{plan['id']}/commit", headers=headers).status_code == 200
        assert service.export_rows(batch)[0]["result"] == "ESCALAR"
        assert "SHA-256" in client.get(f"/sources/audit?batch={batch}").text
    assert service.store.verify_audit()["valid"]


def test_policy_update_requires_confirmation(bundle):
    service, batch = processed(bundle)
    policy = service.store.source(service.batch(batch)["policy_id"])
    policy["version"] = "v3-review"
    app = create_app(service.store.root)
    with TestClient(app) as client:
        headers={"X-CSRF-Token":app.state.service.store.csrf_token}
        args={"files":{"file":("rules.json",json.dumps(policy))},"headers":headers}
        assert client.post(f"/api/batches/{batch}/changes/upload",data={"kind":"policy"},**args).status_code == 400
        assert client.post(f"/api/batches/{batch}/changes/upload",data={"kind":"policy","rules_ack":"true"},**args).status_code == 200


def test_unreadable_history_layout_rejected(bundle):
    _, _, _, path = bundle
    add_archive(path, [("PO-2025-0010",100)])
    book = openpyxl.load_workbook(path)
    book["Pedidos_2025_OLD"]["A1"] = "Campo desconocido"
    book.save(path); book.close()
    with pytest.raises(ValueError,match="columnas Pedido"):
        read_master(path)
