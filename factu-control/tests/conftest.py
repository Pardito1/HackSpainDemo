from pathlib import Path
import json
import pytest
import pymupdf as fitz
import openpyxl

from factu.service import Service
from factu.extract import extract_pdf
from factu.master import read_master

IBAN = (
    "ES4414650100951704302211"  # Deliberately synthetic checksum, as in the challenge.
)


@pytest.fixture(autouse=True)
def sin_credenciales_modelo(monkeypatch):
    """Ningún test toca la red: sin credenciales, leer_campos devuelve sin_clave."""
    for variable in ("CLOUDFLARE_API_TOKEN", "CF_AIG_TOKEN", "LLM_ANTHROPIC_URL"):
        monkeypatch.delenv(variable, raising=False)


def make_pdf(path, total="121,00", iban=IBAN, extra="", raster=False):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        "\n".join(
            [
                "Factura: DEMO-001",
                "Fecha: 01/02/2026",
                "Proveedor: Empresa de prueba",
                "NIF: B12345678",
                "Cliente: NIF B87654321",
                f"IBAN: {iban}",
                "Pedido: PO-2026-0001",
                "Base imponible: 100,00",
                "IVA (21%): 21,00",
                f"TOTAL: {total}",
                extra,
            ]
        ),
        fontsize=12,
    )
    if raster:
        png = page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")
        output = fitz.open()
        output.new_page().insert_image(page.rect, stream=png)
        output.save(path)
        output.close()
    else:
        doc.save(path)
    doc.close()


def make_workbook(path):
    book = openpyxl.Workbook()
    s = book.active
    s.title = "Proveedores"
    s.append(["ProveedorID", "Nombre", "NIF", "IBAN"])
    s.append(["P001", "Empresa de prueba", "B12345678", IBAN])
    s.append(["P001", "Empresa de prueba", "B12345678", IBAN])
    orders = book.create_sheet("Pedidos_2026")
    orders.append(["Pedido", "ProveedorID", "NIF", "Importe", "Estado"])
    orders.append(["PO-2026-0001", "P001", None, 121, "PENDIENTE"])
    rules = book.create_sheet("Norma_Pagos_v3")
    rules.append(["Política sintética de pruebas"])
    stale = book.create_sheet("ANTIGUO_NO_USAR")
    stale.append(["=1/0"])
    book.save(path)
    book.close()


@pytest.fixture
def erp_snapshot():
    return {
        "complete": True,
        "total": 1,
        "pages": [],
        "rows": [
            {
                "id": "1",
                "fecha": "2026-02-01",
                "proveedor": "P001",
                "nif": "B12345678",
                "pedido": "PO-2026-0001",
                "importe": "121.00",
                "estado": "PENDIENTE",
            }
        ],
    }


@pytest.fixture
def bundle(tmp_path, erp_snapshot):
    folder = tmp_path / "pdfs"
    folder.mkdir()
    make_pdf(folder / "factura_ñ.pdf")
    workbook = tmp_path / "maestro.xlsx"
    make_workbook(workbook)
    service = Service(tmp_path / "state")
    batch = service.ingest(folder, workbook, "Test", "2026-09-19")["batch_id"]
    snapshot_id = service.store.put_source("erp", erp_snapshot)
    with service.store.connect() as db:
        db.execute("UPDATE batches SET snapshot_id=? WHERE id=?", (snapshot_id, batch))
    return service, batch, folder, workbook


@pytest.fixture
def facts(bundle, erp_snapshot):
    service, batch, folder, workbook = bundle
    return (
        extract_pdf(folder / "factura_ñ.pdf", ocr=False),
        read_master(workbook),
        erp_snapshot,
        service.store.source(service.batch(batch)["policy_id"]),
    )
