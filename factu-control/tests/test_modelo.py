"""Método de extracción "modelo": todo sin red (leer_campos y httpx fingidos).

El PDF de estos tests es una imagen (raster) y el OCR local se sustituye por
tokens deterministas: así se controla exactamente qué campos quedan sin leer
y se comprueba la fusión OCR/modelo, la validación y la caché.
"""

import copy
import json
from decimal import Decimal

import openpyxl
import pytest

from factu import modelo
from factu.extract import extract_pdf, fusionar_modelo
from factu.policy import evaluate
from factu.service import Service
from conftest import make_pdf

NIF_OK = "B12345674"  # dígito de control correcto (el B12345678 de conftest no)
IBAN_OK = "ES9121000418450200051332"  # checksum mod 97 correcto
IBAN_OK_2 = "ES7921000813610123456789"  # también correcto, distinto del maestro

LINEAS = [
    "Factura: DEMO-001",
    "Fecha: 01/02/2026",
    f"NIF: {NIF_OK}",
    f"IBAN: {IBAN_OK}",
    "Pedido: PO-2026-0001",
    "Base imponible: 100,00",
    "IVA (21%): 21,00",
    "TOTAL: 121,00",
    "Moneda: EUR",
]
USO = {
    "backend": "cf_workers_ai",
    "modelo": "@cf/meta/llama-4-scout-17b-16e-instruct",
    "tokens_in": 2374,
    "tokens_out": 210,
    "neurons": 66.7,
    "segundos": 2.5,
}


def tokens_ocr(lineas):
    return [
        {"text": t, "bbox": [50, 40 + 20 * i, 500, 52 + 20 * i], "confidence": 0.99}
        for i, t in enumerate(lineas)
    ]


def finge_ocr(monkeypatch, lineas):
    monkeypatch.setattr(
        "factu.extract.ocr_page", lambda page, scale=3: tokens_ocr(lineas)
    )


def finge_modelo(monkeypatch, respuesta, contador=None):
    def fake(path_pdf, paginas):
        if contador is not None:
            contador.append((str(path_pdf), list(paginas)))
        return respuesta

    monkeypatch.setattr(modelo, "leer_campos", fake)


def make_master(path):
    book = openpyxl.Workbook()
    s = book.active
    s.title = "Proveedores"
    s.append(["ProveedorID", "Nombre", "NIF", "IBAN"])
    s.append(["P001", "Empresa de prueba", NIF_OK, IBAN_OK])
    orders = book.create_sheet("Pedidos_2026")
    orders.append(["Pedido", "ProveedorID", "NIF", "Importe", "Estado"])
    orders.append(["PO-2026-0001", "P001", None, 121, "PENDIENTE"])
    rules = book.create_sheet("Norma_Pagos_v3")
    rules.append(["Política sintética de pruebas"])
    book.save(path)
    book.close()


@pytest.fixture
def lote_escaneado(tmp_path):
    folder = tmp_path / "pdfs"
    folder.mkdir()
    make_pdf(folder / "escaneada.pdf", raster=True)
    workbook = tmp_path / "maestro.xlsx"
    make_master(workbook)
    service = Service(tmp_path / "state")
    batch = service.ingest(folder, workbook, "Modelo", "2026-09-19")["batch_id"]
    snapshot_id = service.store.put_source(
        "erp",
        {
            "complete": True,
            "total": 1,
            "pages": [],
            "rows": [
                {
                    "id": "1",
                    "fecha": "2026-02-01",
                    "proveedor": "P001",
                    "nif": NIF_OK,
                    "pedido": "PO-2026-0001",
                    "importe": "121.00",
                    "estado": "PENDIENTE",
                }
            ],
        },
    )
    with service.store.connect() as db:
        db.execute("UPDATE batches SET snapshot_id=? WHERE id=?", (snapshot_id, batch))
    return service, batch, folder, workbook


def detalle(service):
    doc = service.store.one("SELECT * FROM documents ORDER BY created,id")
    return service.detail(doc["id"])


# --- (a) el modelo completa lo que el OCR no leyó y decide la política --------


def test_model_fills_missing_field_and_policy_decides(lote_escaneado, monkeypatch):
    service, batch, _, _ = lote_escaneado
    monkeypatch.setenv("LLM_PRECIO_NEURONA_MIL", "0.011")
    finge_ocr(monkeypatch, [l for l in LINEAS if "IBAN" not in l])
    finge_modelo(
        monkeypatch,
        {
            "campos": {
                "iban": {
                    "raw_value": IBAN_OK,
                    "evidencia": f"IBAN: {IBAN_OK}",
                    "page": 1,
                }
            },
            "uso": dict(USO),
        },
    )
    assert service.process(batch, ocr=True)["decisions"] == 1
    assert service.export_rows(batch) == [
        {"file_id": "escaneada.pdf", "result": "PAGAR"}
    ]
    fields = detalle(service)["extraction"]["fields"]
    assert fields["iban"]["status"] == "OK" and fields["iban"]["value"] == IBAN_OK
    prueba = [e for e in fields["iban"]["evidence"] if e["method"] == "modelo"][0]
    assert prueba["source_text"] == f"IBAN: {IBAN_OK}" and prueba["page"] == 1
    assert prueba["transformations"] == ["modelo", "identifier"]
    # paso 4: el coste queda registrado con neuronas y precio de entorno
    row = service.store.one("SELECT * FROM costs WHERE stage='modelo'")
    assert row["external_eur"] == pytest.approx(66.7 / 1000 * 0.011)
    payload = json.loads(row["payload"])
    assert payload["backend"] == "cf_workers_ai"
    assert payload["neurons"] == 66.7 and payload["precio_configurado"] is True


# --- (b) error del proveedor: aviso MODELO_NO_DISPONIBLE y ESCALAR ------------


def test_model_does_not_invent_currency(lote_escaneado, monkeypatch):
    service, batch, _, _ = lote_escaneado
    finge_ocr(monkeypatch, [line for line in LINEAS if "Moneda" not in line and "IBAN" not in line])
    finge_modelo(monkeypatch, {"campos": {"iban": {
        "raw_value": IBAN_OK, "evidencia": f"IBAN: {IBAN_OK}", "page": 1}}, "uso": dict(USO)})
    service.process(batch, ocr=True)
    assert detalle(service)["extraction"]["fields"]["currency"]["status"] == "MISSING"
    assert service.export_rows(batch)[0]["result"] == "ESCALAR"


def test_model_error_warns_and_escalates(lote_escaneado, monkeypatch):
    service, batch, _, _ = lote_escaneado
    finge_ocr(monkeypatch, [l for l in LINEAS if "IBAN" not in l])
    finge_modelo(monkeypatch, {"error": "timeout"})
    service.process(batch, ocr=True)
    assert service.export_rows(batch)[0]["result"] == "ESCALAR"
    info = detalle(service)
    aviso = [
        w
        for w in info["extraction"]["warnings"]
        if w["code"] == "MODELO_NO_DISPONIBLE"
    ][0]
    assert aviso["message"].startswith("timeout") and "iban" in aviso["fields"]
    assert info["extraction"]["fields"]["iban"]["status"] == "MISSING"
    assert service.store.one("SELECT count(*) n FROM costs WHERE stage='modelo'")["n"] == 0


# --- (c) OCR y modelo en desacuerdo: CONFLICT, nunca sustitución ---------------


def test_model_disagreeing_with_ocr_is_conflict(lote_escaneado, monkeypatch):
    service, batch, _, _ = lote_escaneado
    finge_ocr(monkeypatch, [l for l in LINEAS if "TOTAL" not in l])
    finge_modelo(
        monkeypatch,
        {
            "campos": {
                "iban": {
                    "raw_value": IBAN_OK_2,
                    "evidencia": f"IBAN: {IBAN_OK_2}",
                    "page": 1,
                },
                "total": {"raw_value": "121,00", "evidencia": "TOTAL: 121,00", "page": 1},
            },
            "uso": dict(USO),
        },
    )
    service.process(batch, ocr=True)
    assert service.export_rows(batch)[0]["result"] == "ESCALAR"
    fields = detalle(service)["extraction"]["fields"]
    assert fields["iban"]["status"] == "CONFLICT" and fields["iban"]["value"] is None
    assert fields["total"]["status"] == "OK" and fields["total"]["value"] == "121.00"


# --- (d) segunda pasada: caché de extracción, sin nueva llamada ----------------


def test_extraction_cache_skips_second_model_call(lote_escaneado, monkeypatch):
    service, batch, folder, workbook = lote_escaneado
    llamadas = []
    finge_ocr(monkeypatch, [l for l in LINEAS if "IBAN" not in l])
    finge_modelo(
        monkeypatch,
        {
            "campos": {
                "iban": {
                    "raw_value": IBAN_OK,
                    "evidencia": f"IBAN: {IBAN_OK}",
                    "page": 1,
                }
            },
            "uso": dict(USO),
        },
        contador=llamadas,
    )
    service.process(batch, ocr=True)
    second = service.ingest(folder, workbook, "Second", "2026-09-19")["batch_id"]
    with service.store.connect() as db:
        db.execute(
            "UPDATE batches SET snapshot_id=? WHERE id=?",
            (service.batch(batch)["snapshot_id"], second),
        )
    service.process(second, ocr=True)
    assert len(llamadas) == 1
    counts = {
        r["stage"]: r["n"]
        for r in service.store.all(
            "SELECT stage, count(*) n FROM costs WHERE stage IN"
            " ('modelo','extract_cache') GROUP BY stage"
        )
    }
    assert counts == {"modelo": 1, "extract_cache": 1}
    assert service.export_rows(second)[0]["result"] == "NO_PAGAR"  # copia confirmada


# --- (e) sin credenciales: sin_clave y ni un byte a la red ---------------------


def test_missing_credentials_never_touch_network(tmp_path, monkeypatch):
    def boom(*args, **kwargs):
        raise AssertionError("se intentó una llamada de red sin credenciales")

    monkeypatch.setattr(modelo.httpx, "post", boom)
    assert modelo.leer_campos("no-importa.pdf", [1]) == {"error": "sin_clave"}
    path = tmp_path / "escaneada.pdf"
    make_pdf(path, raster=True)
    finge_ocr(monkeypatch, [l for l in LINEAS if "IBAN" not in l])
    extraction = extract_pdf(path, ocr=True)
    aviso = [
        w for w in extraction["warnings"] if w["code"] == "MODELO_NO_DISPONIBLE"
    ][0]
    assert aviso["message"] == "sin_clave"
    assert extraction["engines"]["modelo"]["disponible"] is False


def test_ocr_disabled_never_calls_model(tmp_path, monkeypatch):
    finge_modelo(
        monkeypatch,
        pytest.fail,  # cualquier llamada rompe el test
    )
    path = tmp_path / "escaneada.pdf"
    make_pdf(path, raster=True)
    extraction = extract_pdf(path, ocr=False)
    assert any(w["code"] == "OCR_DISABLED" for w in extraction["warnings"])


# --- (f) validación de lecturas del modelo (casos del ensayo real) -------------


@pytest.mark.parametrize(
    "campo,valor,evidencia,motivo",
    [
        ("supplier_nif", "B963120774", None, "nif_formato"),  # scan_002: 10 caracteres
        ("supplier_nif", "B9623341", None, "nif_formato"),  # scan_023: 8 caracteres
        ("supplier_nif", "A58231074", None, "cif_cliente"),  # scan_021: CIF del cliente
        ("iban", "E393 8888 6371 8888 5555 21", None, "iban_formato"),  # scan_013
        ("order", "PO20260487", None, "formato_pedido"),  # scan_010/016: sin guiones
        ("total", "5310.00", "Total a pagar en plazo", "inferido"),  # calculado, no leído
    ],
)
def test_model_reading_rejected(campo, valor, evidencia, motivo):
    assert modelo.validar_lectura_modelo(campo, valor, evidencia).startswith(motivo)


@pytest.mark.parametrize(
    "campo,valor,evidencia",
    [
        ("supplier_nif", NIF_OK, None),
        ("iban", IBAN_OK, None),
        ("supplier_nif", "A41220987", None),  # del maestro: checksum incorrecto
        ("iban", "ES21 0049 1500 0512 3456 7890", None),  # del maestro: mod 97 falla
        ("order", "PO-2026-0487", None),
        ("total", "5310.00", "TOTAL: 5.310,00"),
        ("total", "535.35", "Total 535,35."),  # puntuación final en la evidencia
        ("tax_rate", "21", "IVA (21%): 1.115,10"),
        ("date", "01/02/2026", None),
    ],
)
def test_model_reading_accepted(campo, valor, evidencia):
    assert modelo.validar_lectura_modelo(campo, valor, evidencia) is None


def test_checksum_es_diagnostico_no_bloqueo():
    """Identificadores del maestro con checksum sintético se aceptan y quedan
    marcados con checksum=False en el candidato; la política hace el contraste."""
    fields = {
        f: {"status": "MISSING", "value": None, "evidence": []}
        for f in ("supplier_nif", "iban", "total")
    }
    fusionar_modelo(
        fields,
        {
            "supplier_nif": {"raw_value": "A41220987", "evidencia": "NIF: A41220987", "page": 1},
            "iban": {
                "raw_value": "ES21 0049 1500 0512 3456 7890",
                "evidencia": "IBAN: ES21 0049 1500 0512 3456 7890",
                "page": 1,
            },
            "total": {"raw_value": "121,00", "evidencia": "TOTAL: 121,00", "page": 1},
        },
    )
    assert fields["supplier_nif"]["status"] == "OK"
    assert fields["supplier_nif"]["evidence"][-1]["checksum"] is False
    assert fields["iban"]["status"] == "OK"
    assert fields["iban"]["evidence"][-1]["checksum"] is False
    assert fields["total"]["evidence"][-1]["checksum"] is None


def test_pedido_solo_del_modelo_no_produce_no_pagar(facts):
    """Un pedido leído únicamente por el modelo con asiento PAGADA se consulta,
    nunca se rechaza; leído por texto nativo u OCR sigue siendo NO_PAGAR."""
    extraction, master, snapshot, policy = facts
    pagado = dict(snapshot, rows=[dict(snapshot["rows"][0], estado="PAGADA")])
    base = evaluate(extraction, master, pagado, policy, "2026-09-19")
    assert base["result"] == "NO_PAGAR"

    solo_modelo = copy.deepcopy(extraction)
    fact = solo_modelo["fields"]["order"]
    fact["evidence"] = [
        {
            "raw_value": fact["value"],
            "value": fact["value"],
            "error": None,
            "page": 1,
            "bbox": None,
            "coordinate_space": None,
            "method": "modelo",
            "confidence": None,
            "source_text": "Pedido: " + fact["value"],
            "checksum": None,
            "transformations": ["modelo", "identifier"],
        }
    ]
    decision = evaluate(solo_modelo, master, pagado, policy, "2026-09-19")
    assert decision["result"] == "ESCALAR"
    assert any(
        "lo leyó el modelo y el ERP lo da por pagado" in q for q in decision["questions"]
    )


# --- (g) parseo de la salida del modelo ----------------------------------------


@pytest.mark.parametrize(
    "texto",
    [
        '{"total":{"value":"5.310,00","evidencia":"TOTAL: 5.310,00","pagina":1}}',
        '```json\n{"total":{"value":"5.310,00","evidencia":"TOTAL: 5.310,00","pagina":1}}\n```',
        '```\n{"total":{"value":"5.310,00","evidencia":"TOTAL: 5.310,00","pagina":1}}\n```',
        'Claro, aquí tienes:\n{"total":{"value":"5.310,00","evidencia":"TOTAL: 5.310,00","pagina":1}} ¡Listo!',
    ],
)
def test_model_output_json_with_or_without_fences(texto):
    assert modelo.extraer_json(texto)["total"]["value"] == "5.310,00"


def test_model_output_number_is_decimal_not_float():
    parsed = modelo.extraer_json('{"total":{"value":395.00}}')
    assert parsed["total"]["value"] == Decimal("395.00")
    assert modelo.normalizar_importe(parsed["total"]["value"]) == Decimal("395.00")


def test_model_output_garbage_is_none():
    assert modelo.extraer_json("no hay json aquí") is None
    assert modelo.extraer_json('{"roto": ') is None
    assert modelo.extraer_json("") is None


@pytest.mark.parametrize(
    "crudo,limpio",
    [
        ("395.00", "395.00"),
        ("938.18", "938.18"),
        ("379,61", "379.61"),
        ("5.310,00", "5310.00"),
    ],
)
def test_amount_formats_normalized(crudo, limpio):
    assert modelo.normalizar_importe(crudo) == Decimal(limpio)


def test_tax_rate_normalized_as_int():
    assert modelo.normalizar_tipo(21) == 21
    assert modelo.normalizar_tipo("21%") == 21
    assert modelo.normalizar_tipo("21 %") == 21
    with pytest.raises(ValueError):
        modelo.normalizar_tipo("veintiuno")


# --- (h) cortacircuito: tres errores de red seguidos y no se llama más ---------


def test_circuit_opens_after_three_network_errors(tmp_path, monkeypatch):
    llamadas = []

    def post_con_timeout(*args, **kwargs):
        llamadas.append(1)
        raise modelo.httpx.TimeoutException("sin respuesta")

    monkeypatch.setenv("CF_ACCOUNT_ID", "cuenta")
    monkeypatch.setenv("CLOUDFLARE_API_TOKEN", "token-de-prueba")
    monkeypatch.setattr(modelo.httpx, "post", post_con_timeout)
    path = tmp_path / "escaneada.pdf"
    make_pdf(path, raster=True)
    modelo.reiniciar_circuito()
    try:
        for _ in range(modelo.UMBRAL_CIRCUITO):
            assert modelo.leer_campos(path, [1]) == {"error": "timeout"}
        previas = len(llamadas)
        assert modelo.leer_campos(path, [1]) == {"error": "circuito_abierto"}
        assert len(llamadas) == previas  # la cuarta lectura no toca la red
        modelo.reiniciar_circuito()  # lo que hace service.process() al arrancar
        assert modelo.leer_campos(path, [1]) == {"error": "timeout"}
        assert len(llamadas) == previas + 1
    finally:
        modelo.reiniciar_circuito()
