import copy
import json

import openpyxl
import pytest

from factu.master import read_master
from factu.presentation import invoice_activity


def test_current_and_historical_columns_by_name(bundle):
    _, _, _, path = bundle
    book = openpyxl.load_workbook(path)
    current = book['Pedidos_2026']
    current.delete_rows(1, current.max_row)
    current.append(['Estado', ' importe_total ', 'NIF', 'Pedido', 'ProveedorID'])
    current.append(['PENDIENTE', 121, 'B12345678', 'PO-2026-0001', 'P001'])
    history = book.create_sheet('Pedidos_2025_OLD')
    history.append(['Pedido', 'ProveedorID', 'NIF', 'Importe_Total', 'Estado', 'Fecha_Pedido'])
    history.append(['PO-2025-9001', 'P101', 'B50123456', 4823.6, 'CERRADO', '2025-01-01'])
    book.save(path); book.close()
    original = path.read_bytes()
    result = read_master(path)
    current = result['orders']['PO-2026-0001'][0]
    assert current['total'] == '121' and current['source']['cells']['total'] == 'B2'
    assert current['source']['cells']['order'] == 'D2'
    old = result['historical']['orders']['PO-2025-9001'][0]
    assert old['total'] == '4823.6' and old['source']['cells']['total'] == 'D2'
    assert old['supplier_id'] == 'P101' and old['state'] == 'CERRADO'
    assert result['historical']['coverage'] == 'partial'
    assert path.read_bytes() == original


@pytest.mark.parametrize('sheet', ['Pedidos_2026', 'Pedidos_2025_OLD'])
def test_ambiguous_amount_headers_rejected(bundle, sheet):
    _, _, _, path = bundle
    book = openpyxl.load_workbook(path)
    tab = book[sheet] if sheet in book.sheetnames else book.create_sheet(sheet)
    if sheet.endswith('OLD'):
        tab.append(['Pedido', 'Importe'])
    tab.cell(1, tab.max_column + 1, 'Importe_Total')
    book.save(path); book.close()
    with pytest.raises(ValueError, match='varias columnas'):
        read_master(path)


def repeated_history(bundle):
    service, batch, *_ = bundle
    service.process(batch, ocr=False)
    doc = service.store.one('SELECT id FROM documents')['id']
    for fingerprint in ('second-version', 'third-version'):
        service.code_sha256 = fingerprint
        service.reevaluate_all()
    return service, service.detail(doc)


def test_automatic_checks_group_without_deleting_audit(bundle):
    service, detail = repeated_history(bundle)
    before = copy.deepcopy(detail)
    view = invoice_activity(detail)
    assert len(view['activity']) == 1
    group = view['activity'][0]
    assert group['title'] == '3 comprobaciones automáticas · mismo resultado'
    assert len(group['check_runs']) == 3
    assert group['check_runs'][-1]['cause'] == 'actualización de la aplicación'
    assert detail == before and len(detail['history']) == 3
    assert service.store.verify_audit()['valid']


def test_same_result_different_field_or_evidence_not_grouped(bundle):
    _, detail = repeated_history(bundle)
    changed = json.loads(detail['history'][0]['payload'])
    changed['fields']['total']['value'] = '122'
    detail['history'][0]['payload'] = json.dumps(changed)
    assert len(invoice_activity(detail)['activity']) == 2
    changed['fields']['total']['value'] = '121.00'
    changed['rules'][0]['state'] = 'FAIL'
    detail['history'][0]['payload'] = json.dumps(changed)
    assert len(invoice_activity(detail)['activity']) >= 2


def test_human_review_breaks_automatic_group(bundle):
    _, detail = repeated_history(bundle)
    detail['events'].insert(1, {'kind': 'human_correction', 'created': '2026-09-19T12:00:00',
                              'payload': {'field': 'iban', 'value': 'ES123', 'actor': 'Marta', 'reason': 'Confirmado en original'}})
    activity = invoice_activity(detail)['activity']
    assert [item['kind'] for item in activity] == ['check', 'review', 'check']
