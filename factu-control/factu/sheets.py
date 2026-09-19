"""Read-only, bounded view of the exact preserved XLSX. No formula evaluation."""
import re
from datetime import date, datetime
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter

from .utils import digest

ROLES = {"Proveedores": "Identidad y cuenta bancaria de los proveedores.",
         "Pedidos_2026": "Pedidos actuales que contrastamos con la contabilidad.",
         "Norma_Pagos_v3": "Norma de referencia para las reglas de decisión.",
         "Pedidos_2025_OLD": "Archivo histórico parcial. Señala coincidencias; no demuestra un pago duplicado."}


def workbook_path(service, source_id):
    source = service.store.one("SELECT kind,payload FROM sources WHERE id=?", (source_id,))
    if not source or source["kind"] != "master":
        raise ValueError("Excel registrado no encontrado.")
    master = service.store.source(source_id)
    sha = master.get("sha256", "")
    if not re.fullmatch(r"[0-9a-f]{64}", sha):
        raise ValueError("No se puede verificar este Excel.")
    # Never accept a user-supplied filesystem path, including the stored blob path.
    path = service.store.root / "blobs" / (sha + ".xlsx")
    if not path.is_file() or digest(path.read_bytes()) != sha:
        raise ValueError("El Excel original no está disponible o ha cambiado. Restaura la copia conservada.")
    return path, master


def sheet_view(service, source_id, sheet_name, page=1, column_page=1):
    path, master = workbook_path(service, source_id)
    with ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Excel demasiado grande para el visor.")
    book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError("Hoja no encontrada en este Excel.")
        sheet = book[sheet_name]
        # Explicit bound avoids a malicious sparse sheet forcing a million-row scan.
        if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 200:
            raise ValueError("Esta hoja supera el tamaño del visor (20.000 filas o 200 columnas). Puedes descargar el original.")
        pages = max(1, ((sheet.max_row or 1) + 49) // 50)
        column_pages = max(1, ((sheet.max_column or 1) + 7) // 8)
        if not 1 <= page <= pages or not 1 <= column_page <= column_pages:
            raise ValueError("Página de la hoja no encontrada.")
        start, end = (page - 1) * 50 + 1, min(page * 50, sheet.max_row or 1)
        first_col, last_col = (column_page - 1) * 8 + 1, min(column_page * 8, sheet.max_column or 1)
        rows = []
        for n, row in enumerate(sheet.iter_rows(min_row=start, max_row=end, min_col=first_col, max_col=last_col), start):
            cells = []
            for cell in row:
                value = cell.value
                if isinstance(value, (date, datetime)):
                    value = value.strftime("%d/%m/%Y")
                cells.append({"value": "" if value is None else str(value), "formula": cell.data_type == "f"})
            rows.append({"number": n, "cells": cells})
        return {"master": master, "source_id": source_id, "sheet_name": sheet_name, "sheet_names": book.sheetnames,
                "role": ROLES.get(sheet_name, "Hoja conservada como referencia. No se utiliza para decidir."),
                "rows": rows, "columns": [get_column_letter(c) for c in range(first_col, last_col + 1)],
                "page": page, "pages": pages, "column_page": column_page, "column_pages": column_pages,
                "start": start, "end": end, "row_count": sheet.max_row or 1}
    finally:
        book.close()
