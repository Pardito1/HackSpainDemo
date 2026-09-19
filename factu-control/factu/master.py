from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path
import openpyxl
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from .utils import clean, digest, iban_checksum, identifier, invoice_date, money
from .historical import read_history
from .columns import ORDER_COLUMNS, resolve_columns


LOTE2_SUPPLIER_COLUMNS = ("ID", "Razon Social", "NIF", "IBAN", "Ciudad", "Condiciones")
LOTE2_ORDER_COLUMNS = (
    "pedido",
    "proveedor_id",
    "nif",
    "importe_total",
    "estado",
    "fecha_pedido",
)


def read_master(path):
    try:
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Excel demasiado grande una vez descomprimido")
        return _read_master(path)
    except (BadZipFile, KeyError, IndexError, TypeError, ParseError, SyntaxError,
            openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise ValueError("Excel dañado o con estructura inválida. Adjunta un archivo .xlsx válido con las hojas requeridas.") from exc


def order_date(value):
    """Fecha de un pedido del maestro; None si la celda está vacía o ilegible."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    if not text:
        return None
    try:
        return invoice_date(text)
    except ValueError:
        return None


def _read_lote2_csv(path, expected_columns, label):
    """Read one official incremental CSV without guessing its schema.

    The rows are kept in the composed master with their exact file hash and
    row/header provenance.  They are not copied into the workbook or silently
    turned into its rows: the original Excel remains viewable as received.
    """
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"No se encuentra {label}: {path}")
    raw = path.read_bytes()
    if not raw or len(raw) > 10 * 1024 * 1024:
        raise ValueError(f"{label} está vacío o supera el límite de 10 MB")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(f"{label} debe estar codificado en UTF-8") from exc
    reader = csv.DictReader(io.StringIO(text, newline=""))
    headers = reader.fieldnames or []
    if len(headers) != len(set(headers)) or set(headers) != set(expected_columns):
        expected = ", ".join(expected_columns)
        raise ValueError(f"{label}: se esperan exactamente las columnas {expected}")
    rows = list(reader)
    if not rows:
        raise ValueError(f"{label} no contiene filas")
    return raw, rows, headers


def read_lote2_master(workbook, suppliers_csv, orders_csv, store_blob=None):
    """Compose the Lote 2 master while preserving every source separately.

    ``FINAL_v7_DEFINITIVO_ahorasi.xlsx`` remains the canonical workbook.
    ``proveedores_nuevos.csv`` and ``pedidos_nuevos.csv`` are incremental
    sources with their own hash, blob and per-row provenance.  A collision is
    retained as multiple source rows so the policy can escalate ambiguity;
    nothing from the CSV files overwrites a previous record in place.
    """
    workbook = Path(workbook)
    master = read_master(workbook)
    suppliers_raw, supplier_rows, supplier_headers = _read_lote2_csv(
        suppliers_csv, LOTE2_SUPPLIER_COLUMNS, "proveedores_nuevos.csv"
    )
    orders_raw, order_rows, order_headers = _read_lote2_csv(
        orders_csv, LOTE2_ORDER_COLUMNS, "pedidos_nuevos.csv"
    )

    def source_file(kind, path, raw, rows, headers, role):
        path = Path(path)
        payload = {
            "kind": kind,
            "filename": path.name,
            "sha256": digest(raw),
            "rows": len(rows),
            "columns": list(headers),
            "role": role,
        }
        if store_blob is not None:
            payload["blob"] = str(store_blob(raw, ".csv"))
        return payload

    supplier_file = source_file(
        "lote2_providers_csv",
        suppliers_csv,
        suppliers_raw,
        supplier_rows,
        supplier_headers,
        "Altas o cambios de identidad y cuenta de proveedores para Lote 2.",
    )
    order_file = source_file(
        "lote2_orders_csv",
        orders_csv,
        orders_raw,
        order_rows,
        order_headers,
        "Pedidos incrementales de Lote 2 para contrastar con el ERP.",
    )

    def row_source(file_meta, row_number, cells):
        return {
            "kind": "csv",
            "filename": file_meta["filename"],
            "sha256": file_meta["sha256"],
            "row": row_number,
            "cells": {name: f"{name}#{row_number}" for name in cells},
        }

    for row_number, row in enumerate(supplier_rows, start=2):
        supplier_id = clean(row["ID"])
        nif, iban = identifier(row["NIF"]), identifier(row["IBAN"])
        if not supplier_id or not clean(row["Razon Social"]) or not nif or not iban:
            raise ValueError(f"proveedores_nuevos.csv: fila {row_number} incompleta")
        supplier = {
            "id": supplier_id,
            "name": clean(row["Razon Social"]),
            "nif": nif,
            "iban": iban,
            "city": clean(row["Ciudad"]),
            "conditions": clean(row["Condiciones"]),
            "source": row_source(supplier_file, row_number, LOTE2_SUPPLIER_COLUMNS),
        }
        supplier["iban_checksum_diagnostic"] = iban_checksum(supplier["iban"])
        # Do not replace an existing source row: source identity conflicts are
        # policy evidence, not an import-time choice.
        master["suppliers"].setdefault(supplier["id"], []).append(supplier)

    for row_number, row in enumerate(order_rows, start=2):
        order_id = identifier(row["pedido"])
        supplier_id, nif = clean(row["proveedor_id"]), identifier(row["nif"])
        if not order_id or not supplier_id or not nif:
            raise ValueError(f"pedidos_nuevos.csv: fila {row_number} incompleta")
        try:
            total = str(money(row["importe_total"]))
            order_date = date.fromisoformat(clean(row["fecha_pedido"])).isoformat()
        except (ValueError, ArithmeticError) as exc:
            raise ValueError(f"pedidos_nuevos.csv: fila {row_number} tiene importe o fecha inválidos") from exc
        order = {
            "id": order_id,
            "supplier_id": supplier_id,
            "nif": nif,
            "total": total,
            "state": clean(row["estado"]),
            "date": order_date,
            "source": row_source(order_file, row_number, LOTE2_ORDER_COLUMNS),
        }
        master["orders"].setdefault(order["id"], []).append(order)

    master["supplemental_sources"] = [supplier_file, order_file]
    master["lote2_composed"] = {
        "base_excel": {"filename": master["filename"], "sha256": master["sha256"]},
        "suppliers_added": len(supplier_rows),
        "orders_added": len(order_rows),
    }
    return master


def _read_master(path):
    book = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        for name in ("Proveedores", "Pedidos_2026", "Norma_Pagos_v3"):
            if name not in book.sheetnames:
                raise ValueError(
                    f"Falta hoja {name}; hay que configurar un adaptador para este maestro"
                )
        suppliers = {}
        for row in book["Proveedores"].iter_rows(min_row=2, max_col=4):
            if not row[0].value:
                continue
            supplier = {
                "id": clean(row[0].value),
                "name": clean(row[1].value),
                "nif": identifier(row[2].value),
                "iban": identifier(row[3].value),
                "source": {
                    "sheet": "Proveedores",
                    "row": row[0].row,
                    "cells": {
                        "id": f"A{row[0].row}",
                        "name": f"B{row[0].row}",
                        "nif": f"C{row[0].row}",
                        "iban": f"D{row[0].row}",
                    },
                },
            }
            supplier["iban_checksum_diagnostic"] = iban_checksum(supplier["iban"])
            suppliers.setdefault(supplier["id"], []).append(supplier)
        orders, order_dates = {}, []
        columns = resolve_columns(book["Pedidos_2026"], ORDER_COLUMNS,
                                  ("order", "supplier_id", "nif", "total", "state"))
        for source_row in book["Pedidos_2026"].iter_rows(min_row=2):
            row = [source_row[columns[key]] for key in ("order", "supplier_id", "nif", "total", "state")]
            if not row[0].value:
                continue
            order = {
                "id": identifier(row[0].value),
                "supplier_id": clean(row[1].value),
                "nif": identifier(row[2].value),
                "total": str(money(row[3].value)),
                "state": clean(row[4].value),
                "source": {
                    "sheet": "Pedidos_2026",
                    "row": row[0].row,
                    "cells": {key: f"{openpyxl.utils.get_column_letter(columns[key] + 1)}{row[0].row}"
                              for key in ("order", "supplier_id", "nif", "total")},
                },
            }
            orders.setdefault(order["id"], []).append(order)
            if "date" in columns:
                issued = order_date(source_row[columns["date"]].value)
                if issued:
                    order_dates.append(issued)
        rules = [
            {"cell": cell.coordinate, "text": clean(cell.value)}
            for row in book["Norma_Pagos_v3"]
            for cell in row
            if cell.value
        ]
        return {
            "suppliers": suppliers,
            "orders": orders,
            "last_order_date": max(order_dates) if order_dates else None,
            "historical": read_history(book),
            "rules": rules,
            "sha256": digest(path.read_bytes()),
            "filename": path.name,
            "sheets": book.sheetnames,
        }
    finally:
        book.close()
