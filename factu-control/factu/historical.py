"""A partial order archive is evidence for review, never proof of payment."""
import re
from openpyxl.utils import get_column_letter

from .utils import clean, identifier, money
from .columns import ORDER_COLUMNS, resolve_columns

SHEET = "Pedidos_2025_OLD"


def read_history(book):
    result = {"sheet": SHEET, "available": SHEET in book.sheetnames,
              "coverage": "partial", "orders": {}, "notes": [], "warnings": []}
    if not result["available"]:
        return result
    sheet = book[SHEET]
    columns = resolve_columns(sheet, ORDER_COLUMNS, ("order", "total"))
    for row in sheet.iter_rows(min_row=2):
        cells = {key: row[index] for key, index in columns.items()}
        raw = clean(cells["order"].value)
        if not raw:
            continue
        source = {"sheet": SHEET, "row": cells["order"].row,
                  "cells": {key: f"{get_column_letter(index + 1)}{cells['order'].row}" for key, index in columns.items()}}
        key = identifier(raw)
        if cells["order"].data_type == "f" or not re.fullmatch(r"PO-\d{4}-[A-Z0-9-]+", key):
            result["notes"].append({"text": raw, "source": source})
            continue
        try:
            total = str(money(cells["total"].value)) if cells["total"].data_type != "f" else None
        except (ValueError, ArithmeticError, TypeError):
            total = None
        if total is None:
            result["warnings"].append({"order": key, "reason": "Importe histórico no verificable", "source": source})
        record = {"id": key, "total": total, "source": source}
        for field in ("supplier_id", "nif", "state", "date"):
            if field in cells:
                record[field] = clean(cells[field].value) if cells[field].data_type != "f" else None
        result["orders"].setdefault(key, []).append(record)
    return result


def history_dependency(master, order):
    """Only matching rows invalidate a decision; an unrelated archive row does not."""
    history = master.get("historical")
    loaded = history is not None or SHEET not in master.get("sheets", [])
    return {"loaded": loaded, "available": bool(history and history["available"]),
            "coverage": "partial", "rows": (history or {}).get("orders", {}).get(identifier(order), []) if order else []}
