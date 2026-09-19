"""
Mira una factura con tus propios ojos: lo que pone el PDF y lo que saca P1.

    python ver.py                          una factura de ejemplo
    python ver.py factura_1936.pdf         una concreta
    python ver.py --trampas                las 7 con instrucciones escondidas
    python ver.py --raras                  las que no cuadran con el Excel
    python ver.py --lista                  nombres de las 500

Pensado para comprobar a mano, sin fiarse de ningun resumen.
"""

from __future__ import annotations

import logging
import sys
import warnings
from decimal import Decimal
from pathlib import Path

warnings.filterwarnings("ignore")
logging.getLogger("pypdf").setLevel(logging.ERROR)

RAIZ = Path(__file__).resolve().parent
sys.path.insert(0, str(RAIZ))

FACTURAS = RAIZ / "500-sombras-de-alberto-main" / "facturas"
EXCEL = RAIZ / "500-sombras-de-alberto-main" / "FINAL_v7_DEFINITIVO_ahorasi.xlsx"

TRAMPAS = [
    "factura_1936.pdf",
    "2026-07-08_P010.pdf",
    "F26-9007_catering.pdf",
    "F26-8801_suministros.pdf",
    "2026-07-01_P009.pdf",
    "2026-23904_construcciones.pdf",
    "FA-3388_ofimática.pdf",
]

RARAS = [
    "FA-4488_transportes.pdf",      # importe escondido con caracteres invisibles
    "F26-3011_suministros.pdf",     # IBAN escondido igual
    "FA-1123_construcciones.pdf",   # fecha 2026-02-30, que no existe
    "F26-9012_electricidad.pdf",    # IVA al 16%, que no existe en Espana
    "2026-0811-B_catering.pdf",     # el total no cuadra con base + IVA
    "FA-5633_transportes.pdf",      # IBAN distinto del maestro
]


def cargar_maestro():
    try:
        import openpyxl
    except ImportError:
        return {}, {}
    if not EXCEL.exists():
        return {}, {}
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    prov = {}
    for f in wb["Proveedores"].iter_rows(values_only=True):
        if f[0] and f[0] != "ID":
            nif = str(f[2]).strip().upper()
            prov.setdefault(nif, {"nombre": str(f[1]).strip(),
                                  "iban": str(f[3]).replace(" ", "").upper()})
    ped = {str(f[0]).strip(): Decimal(str(f[3]))
           for f in wb["Pedidos_2026"].iter_rows(values_only=True)
           if f[0] and str(f[0]).startswith("PO-")}
    return prov, ped


def texto_del_pdf(ruta: Path) -> str:
    from pypdf import PdfReader

    try:
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(ruta)).pages)
    except Exception as e:
        return f"(no se pudo leer: {e})"


def ver(nombre: str, prov: dict, ped: dict) -> None:
    from pipeline.extraccion import extraer_campos

    ruta = FACTURAS / nombre
    if not ruta.exists():
        print(f"  no existe: {nombre}")
        return

    print()
    print("=" * 74)
    print(f" {nombre}")
    print("=" * 74)

    # --- lo que pone el papel -------------------------------------------------
    crudo = texto_del_pdf(ruta)
    print()
    print(" LO QUE PONE EL PDF")
    print(" " + "-" * 72)
    if len(crudo.strip()) < 50:
        print("   (es una FOTO escaneada: no tiene texto dentro)")
    else:
        for linea in crudo.strip().splitlines()[:22]:
            # se marcan los caracteres invisibles para que se vean
            visible = (linea.replace("​", "[·]").replace(" ", "[ ]")
                            .replace("­", "[-]"))
            print("   " + visible[:70])

    # --- lo que saca P1 -------------------------------------------------------
    c = extraer_campos(ruta, nombre)
    print()
    print(" LO QUE SACA P1")
    print(" " + "-" * 72)
    print(f"   nif      {c.nif_proveedor}")
    print(f"   iban     {c.iban}")
    print(f"   pedido   {c.pedido}")
    print(f"   base     {c.base_imponible}")
    print(f"   iva      {c.iva}")
    print(f"   total    {c.importe}")
    print(f"   fecha    {c.fecha}" + ("" if c.fecha_valida else "   <- NO EXISTE"))
    print(f"   factura  {c.numero_factura}")
    print()
    print(f"   cuadra base+IVA=total : {c.cuadra_aritmetica}")
    print(f"   leida con vision      : {c.uso_ocr}")
    if c.confianza_ocr is not None:
        print(f"   confianza             : {c.confianza_ocr}")
    if c.errores_extraccion:
        for e in c.errores_extraccion:
            print(f"   aviso: {e}")

    if c.senales_riesgo.indicios_inyeccion:
        print()
        print("   *** ESTE DOCUMENTO INTENTA DAR ORDENES A LA IA ***")
        print(f"   {c.senales_riesgo.detalle}")

    # --- que dice el Excel ----------------------------------------------------
    if prov or ped:
        print()
        print(" QUE DICE EL EXCEL")
        print(" " + "-" * 72)
        if c.nif_proveedor in prov:
            p = prov[c.nif_proveedor]
            print(f"   proveedor       {p['nombre']}")
            if c.iban:
                igual = c.iban == p["iban"]
                print(f"   IBAN del maestro {p['iban']}")
                print(f"   IBAN coincide   {'SI' if igual else 'NO  <- OJO, FRAUDE'}")
        else:
            print(f"   el NIF {c.nif_proveedor} NO esta en el maestro")

        if c.pedido in ped:
            esperado = ped[c.pedido]
            print(f"   importe pedido  {esperado}")
            if c.importe is not None:
                dif = Decimal(str(c.importe)) - esperado
                if abs(dif) <= Decimal("0.01"):
                    print("   importe coincide  SI")
                else:
                    print(f"   importe coincide  NO   diferencia {dif:+}")
        elif c.pedido:
            print(f"   el pedido {c.pedido} NO existe en Pedidos_2026")


def main() -> int:
    if not FACTURAS.exists():
        print(f"No encuentro las facturas en {FACTURAS}")
        return 1

    args = sys.argv[1:]
    prov, ped = cargar_maestro()

    if args and args[0] == "--lista":
        nombres = sorted(p.name for p in FACTURAS.glob("*.pdf"))
        for i in range(0, len(nombres), 3):
            print("   " + "".join(f"{n:34}" for n in nombres[i:i + 3]))
        print(f"\n   {len(nombres)} facturas")
        return 0

    if args and args[0] == "--trampas":
        print("\n  LAS 7 FACTURAS QUE INTENTAN ENGANAR A LA IA\n")
        for n in TRAMPAS:
            ver(n, prov, ped)
        return 0

    if args and args[0] == "--raras":
        print("\n  CASOS RAROS QUE MERECE LA PENA MIRAR\n")
        for n in RARAS:
            ver(n, prov, ped)
        return 0

    nombres = args or ["2026-01-08_P001.pdf"]
    for n in nombres:
        ver(n, prov, ped)
    return 0


if __name__ == "__main__":
    sys.exit(main())
