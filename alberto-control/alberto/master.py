from __future__ import annotations

import openpyxl
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from .utils import clean, digest, iban_checksum, identifier, money


def read_master(path):
    try:
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Excel demasiado grande una vez descomprimido")
        return _read_master(path)
    except (BadZipFile, KeyError, IndexError, TypeError, ParseError, SyntaxError,
            openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise ValueError("Excel dañado o con estructura inválida. Adjunta un archivo .xlsx válido con las hojas requeridas.") from exc


def _read_master(path):
    book = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        for name in ("Proveedores", "Pedidos_2026", "Norma_Pagos_v3"):
            if name not in book.sheetnames:
                raise ValueError(
                    f"Falta hoja {name}; hay que configurar un adaptador para este maestro"
                )
        suppliers = {}
        for row in book["Proveedores"].iter_rows(min_row=2, max_col=4):
            if not row[0].value:
                continue
            supplier = {
                "id": clean(row[0].value),
                "name": clean(row[1].value),
                "nif": identifier(row[2].value),
                "iban": identifier(row[3].value),
                "source": {
                    "sheet": "Proveedores",
                    "row": row[0].row,
                    "cells": {
                        "id": f"A{row[0].row}",
                        "name": f"B{row[0].row}",
                        "nif": f"C{row[0].row}",
                        "iban": f"D{row[0].row}",
                    },
                },
            }
            supplier["iban_checksum_diagnostic"] = iban_checksum(supplier["iban"])
            suppliers.setdefault(supplier["id"], []).append(supplier)
        orders = {}
        for row in book["Pedidos_2026"].iter_rows(min_row=2, max_col=5):
            if not row[0].value:
                continue
            order = {
                "id": identifier(row[0].value),
                "supplier_id": clean(row[1].value),
                "nif": identifier(row[2].value),
                "total": str(money(row[3].value)),
                "state": clean(row[4].value),
                "source": {
                    "sheet": "Pedidos_2026",
                    "row": row[0].row,
                    "cells": {
                        "order": f"A{row[0].row}",
                        "supplier_id": f"B{row[0].row}",
                        "nif": f"C{row[0].row}",
                        "total": f"D{row[0].row}",
                    },
                },
            }
            orders.setdefault(order["id"], []).append(order)
        rules = [
            {"cell": cell.coordinate, "text": clean(cell.value)}
            for row in book["Norma_Pagos_v3"]
            for cell in row
            if cell.value
        ]
        return {
            "suppliers": suppliers,
            "orders": orders,
            "rules": rules,
            "sha256": digest(path.read_bytes()),
            "filename": path.name,
            "sheets": book.sheetnames,
        }
    finally:
        book.close()
